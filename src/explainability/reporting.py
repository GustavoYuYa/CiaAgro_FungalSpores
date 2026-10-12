"""Side-by-side PNGs and Excel summaries, following the project export pattern."""

from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw
from pytorch_grad_cam.utils.image import show_cam_on_image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


METRICS = ("energy_inside", "iou", "deletion_auc")


def save_heatmap(path, original_rgb, cam, title, subtitle, panel_size=320):
    """Two panels: original clean reference and CAM over the same reference.

    For Noise the CAM was computed on the noisy input, but is displayed on
    its paired original image. The condition is explicitly included in title.
    Raw CAM arrays are saved separately, before visualization/quantization.
    """
    original = original_rgb.permute(1, 2, 0).cpu().numpy().astype(np.float32)
    overlay = show_cam_on_image(original, cam, use_rgb=True, image_weight=0.5)
    left = Image.fromarray(np.rint(original * 255).astype(np.uint8))
    right = Image.fromarray(overlay)
    canvas = Image.new("RGB", (2 * panel_size + 30, panel_size + 95), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((10, 8), title, fill="black")
    draw.text((10, 26), subtitle, fill="black")
    draw.text((10, 50), "Original (Clean)", fill="black")
    draw.text((panel_size + 20, 50), "CAM sobre a original", fill="black")
    canvas.paste(left.resize((panel_size, panel_size), Image.Resampling.NEAREST), (10, 75))
    canvas.paste(right.resize((panel_size, panel_size), Image.Resampling.NEAREST), (panel_size + 20, 75))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(path)


def summarize(records, group_columns):
    """Each model/condition is summarized separately; std is sample std (ddof=1)."""
    rows = []
    for keys, group in records.groupby(group_columns, dropna=False, sort=True):
        if not isinstance(keys, tuple):
            keys = (keys,)
        identifiers = dict(zip(group_columns, keys))
        for metric in METRICS:
            values = group[metric]
            rows.append({
                **identifiers, "metric": metric, "mean": values.mean(),
                "std": values.std(ddof=1), "valid_count": int(values.count()),
                "total_count": len(group), "missing_count": int(values.isna().sum()),
            })
    return pd.DataFrame(rows)


def export_excel(records, metadata, path):
    """One architecture + one CAM method per workbook, never pooled methods."""

    if records.empty:
        raise ValueError("Cannot export an empty XAI evaluation.")
    for column in ("architecture", "method", "target_policy"):
        if records[column].nunique() != 1:
            raise ValueError(f"Do not mix {column} values in the same workbook.")
    groups = ["model", "condition"]
    tables = {
        "Summary": summarize(records, groups),
        "By_Class": summarize(records, groups + ["true_class", "class_name"]),
        "By_Correctness": summarize(records, groups + ["correct"]),
        "Clean": records[records["condition"] == "Clean"],
        "Noise": records[records["condition"] == "Noise"],
        "Protocol": pd.DataFrame(metadata, columns=["model", "parameter", "value"]),
    }
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "C2" if name != "Protocol" else "B2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            for cell in ws[1]:
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="234E70")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for col, column in enumerate(table.columns, 1):
                width = min(55, max(14, len(str(column)) + 2))
                if column in ("value", "clean_path", "segmented_path", "heatmap_path"):
                    width = 55
                ws.column_dimensions[get_column_letter(col)].width = width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Calibri", size=11)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000000"

"""Run with: python -m unittest discover -s tests -p test_explainability.py"""

from pathlib import Path
import tempfile
import unittest

import numpy as np
import torch
from PIL import Image

from src.datasets.spore_dataset import EsporosDataset, EsporosDatasetNoise
from src.datasets.transforms import base_transform, transform_wout_normalize
from src.explainability.cam import generate_cam
from src.explainability.data import PairedTestDataset, normalization_tensors
from src.explainability.metrics import deletion_auc, deletion_baseline, localization_metrics
from src.explainability.models import cam_layer, reshape_tokens
from src.models.cnn import CNN
from src.models.simplevit import SimpleViT


def make_fixture(root):
    """Non-square source PNGs exercise resize-before-mask semantics."""
    rng = np.random.RandomState(3)
    clean_files, segmented_files = [], []
    for i in range(2):
        original = rng.randint(30, 240, (77, 83, 3), dtype=np.uint8)
        segmented = np.zeros_like(original)
        segmented[20:59, 21:62] = original[20:59, 21:62]
        for folder, pixels, files in (
            ("Normal", original, clean_files),
            ("Segmented", segmented, segmented_files),
        ):
            relative = Path("filtered_cropped_images") / folder / f"class_{i}" / f"image_{i}.png"
            (root / relative).parent.mkdir(parents=True, exist_ok=True)
            Image.fromarray(pixels).save(root / relative)
            files.append(relative.as_posix())
    for condition, files in (("normal", clean_files), ("segmented_black", segmented_files)):
        folder = root / "splits" / condition
        folder.mkdir(parents=True)
        np.save(folder / "test_files.npy", np.array(files))
        np.save(folder / "test_labels.npy", np.array([0, 1], dtype=np.int64))


class DataTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        make_fixture(self.root)

    def tearDown(self):
        self.temporary.cleanup()

    def test_matches_existing_evaluation_exactly(self):
        data = PairedTestDataset(self.root)
        mean = [0.7, 0.6, 0.5]
        std = [0.2, 0.3, 0.4]
        clean = EsporosDataset(data.clean_files, data.labels, self.root, base_transform(mean, std))
        noise = EsporosDatasetNoise(data.segmented_files, data.labels, self.root,
                                   transform_wout_normalize(), True, False, mean, std)
        mean_t, std_t = normalization_tensors({"normalization": {"mean": mean, "std": std}})
        for i in (1, 0):  # Out-of-order access must keep the original Noise seed.
            sample = data[i]
            torch.testing.assert_close((sample["Clean"] - mean_t) / std_t, clean[i][0], rtol=0, atol=0)
            torch.testing.assert_close((sample["Noise"] - mean_t) / std_t, noise[i][0], rtol=0, atol=0)
            segmented, _ = data.segmented[i]
            expected_mask = segmented.max(dim=0, keepdims=True).values > 0
            np.testing.assert_array_equal(sample["mask"], expected_mask[0].numpy())
            expected_noise = np.random.RandomState(i).normal(0.5, 0.2, tuple(segmented.shape))
            expected_noise = torch.from_numpy(np.clip(expected_noise, 0, 1).astype(np.float32))
            expected_rgb = segmented * expected_mask + expected_noise * (~expected_mask)
            torch.testing.assert_close(sample["Noise"], expected_rgb, rtol=0, atol=0)

    def test_rejects_misaligned_pairs(self):
        path = self.root / "splits" / "segmented_black" / "test_files.npy"
        np.save(path, np.load(path)[::-1])
        with self.assertRaisesRegex(ValueError, "Unpaired"):
            PairedTestDataset(self.root)


class MetricTests(unittest.TestCase):
    def test_gaussian_deletion_is_repeatable_and_independent_of_global_rng(self):
        rgb = torch.zeros(3, 4, 4)
        mean = torch.ones(3, 1, 1) * 0.5
        first = deletion_baseline(rgb, mean, "gaussian", seed=8)
        np.random.seed(900)
        second = deletion_baseline(rgb, mean, "gaussian", seed=8)
        torch.testing.assert_close(first, second, rtol=0, atol=0)
        self.assertTrue(bool(((first >= 0) & (first <= 1)).all()))

    def test_notebook_localization_formulas(self):
        cam = np.array([[0., 1.], [2., 3.]])
        mask = np.array([[False, False], [True, True]])
        scores = localization_metrics(cam, mask, percentile=50)
        self.assertAlmostEqual(scores["energy_inside"], 5 / 6)
        self.assertEqual(scores["iou"], 1)
        tied = localization_metrics(np.ones((2, 2)), mask)
        self.assertEqual(tied["cam_selected_fraction"], 1)
        self.assertEqual(tied["iou"], 0.5)

    def test_empty_and_zero_maps_are_explicit(self):
        scores = localization_metrics(np.zeros((2, 2)), np.ones((2, 2)))
        self.assertEqual(scores["energy_inside"], 0)
        self.assertEqual(scores["iou"], 0)
        self.assertTrue(scores["cam_zero"])
        empty = localization_metrics(np.ones((2, 2)), np.zeros((2, 2)))
        self.assertFalse(empty["mask_valid"])
        self.assertTrue(np.isnan(empty["energy_inside"]))
        self.assertTrue(np.isnan(empty["iou"]))

    def test_deletion_order_fixed_target_endpoints_and_normalization(self):
        class PixelModel(torch.nn.Module):
            def forward(self, x):
                # Undo normalization so analytical expectations are in RGB space.
                rgb = x * 0.25 + 0.5
                signal = 2 * rgb[:, :, 0, 0].mean(1) - 1
                return torch.stack((signal, -signal), dim=1)

        model = PixelModel().eval()
        rgb = torch.ones(3, 1, 2)
        mean, std = torch.full((3, 1, 1), 0.5), torch.full((3, 1, 1), 0.25)
        cam = np.array([[1., 0.]])
        auc, fractions, probs = deletion_auc(model, rgb, cam, 0, mean, std,
                                             steps=2, batch_size=2, baseline="black")
        high = 1 / (1 + np.exp(-2))
        low = 1 - high
        np.testing.assert_allclose(fractions, [0, 0.5, 1])
        np.testing.assert_allclose(probs, [high, low, low], rtol=1e-6)
        self.assertAlmostEqual(auc, 0.25 * (high + low) + 0.5 * low, places=6)
        reverse, _, _ = deletion_auc(model, rgb, cam[:, ::-1].copy(), 0, mean, std,
                                     steps=2, batch_size=1, baseline="black")
        self.assertLess(auc, reverse)
        mean_auc, _, mean_probs = deletion_auc(model, rgb, cam, 0, mean, std,
                                               steps=7, batch_size=3, baseline="mean")
        self.assertAlmostEqual(mean_probs[-1], 0.5)
        self.assertTrue(0 <= mean_auc <= 1)


class CamTests(unittest.TestCase):
    def test_both_methods_both_architectures_and_hook_cleanup(self):
        torch.manual_seed(9)
        configs = {
            "cnn": {"model": {"image_size": 64}},
            "simplevit": {"model": {"image_size": 64, "patch_size": 8}},
        }
        models = {
            "cnn": CNN(channels=(4, 8, 12, 16)).eval(),
            "simplevit": SimpleViT(image_size=64, patch_size=8, num_classes=3,
                                   dim=16, depth=2, heads=2, dim_head=8, mlp_dim=32).eval(),
        }
        x = torch.rand(1, 3, 64, 64)
        for architecture, model in models.items():
            layer, reshape, _ = cam_layer(model, architecture, configs[architecture])
            original_hooks = len(layer._forward_hooks)
            for method in ("EigenCAM", "EigenGradCAM"):
                with self.subTest(architecture=architecture, method=method):
                    result = generate_cam(model, layer, reshape, x, 0, method)
                    self.assertEqual(result.shape, (64, 64))
                    self.assertTrue(np.isfinite(result).all())
                    self.assertGreater(float(result.max()), 0)
                    self.assertEqual(len(layer._forward_hooks), original_hooks)
            first = generate_cam(model, layer, reshape, x, 0, "EigenCAM")
            second = generate_cam(model, layer, reshape, x, 2, "EigenCAM")
            np.testing.assert_allclose(first, second)

    def test_simplevit_keeps_first_token(self):
        tokens = torch.arange(24).reshape(1, 6, 4)
        mapped = reshape_tokens(tokens, 2, 3)
        torch.testing.assert_close(mapped[0, :, 0, 0], tokens[0, 0])
        torch.testing.assert_close(mapped[0, :, 1, 2], tokens[0, -1])


class ReportTests(unittest.TestCase):
    def test_separate_conditions_counts_and_excel_values(self):
        import pandas as pd
        from openpyxl import load_workbook
        from src.explainability.reporting import export_excel

        records = pd.DataFrame([
            {"architecture": "cnn", "method": "EigenCAM", "model": "M1",
             "condition": condition, "true_class": 0, "class_name": "spore",
             "correct": correct, "target_policy": "true", "energy_inside": value,
             "iou": value, "deletion_auc": 0.25}
            for condition, correct, value in (
                ("Clean", True, 0.8), ("Clean", False, 0.4), ("Noise", False, float("nan")),
            )
        ])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "metrics.xlsx"
            export_excel(records, [{"model": "M1", "parameter": "target_policy", "value": "true"}], path)
            workbook = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Summary", "By_Class", "By_Correctness", "Clean", "Noise", "Protocol"])
            rows = list(workbook["Summary"].values)
            clean = dict(zip(rows[0], rows[1]))
            self.assertEqual(clean["condition"], "Clean")
            self.assertAlmostEqual(clean["mean"], 0.6)
            self.assertAlmostEqual(clean["std"], np.std([0.8, 0.4], ddof=1))
            self.assertEqual(clean["valid_count"], 2)
            noise = dict(zip(rows[0], rows[4]))
            self.assertIsNone(noise["mean"])
            self.assertEqual(noise["valid_count"], 0)
            self.assertEqual(noise["missing_count"], 1)
            workbook.close()
            mixed = records.copy()
            mixed.loc[0, "method"] = "EigenGradCAM"
            with self.assertRaisesRegex(ValueError, "Do not mix method"):
                export_excel(mixed, [], path)


if __name__ == "__main__":
    torch.set_num_threads(1)
    unittest.main()
